import email
import io
import logging
import re
import zipfile
from typing import List, Optional

from pypdf import PdfReader
from docx import Document
import icalendar

# Optional dependency imports with safe fallbacks
try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    import xlrd
except ImportError:
    xlrd = None

try:
    from pptx import Presentation
except ImportError:
    Presentation = None

try:
    from striprtf.striprtf import rtf_to_text
except ImportError:
    rtf_to_text = None

try:
    from bs4 import BeautifulSoup
except ImportError:
    BeautifulSoup = None

try:
    import extract_msg
except ImportError:
    extract_msg = None

from extractors.ocr_extractor import extract_text_from_image

logger = logging.getLogger("TextExtractor")

# Safety extraction limits to maintain < 100 MB RAM budget
MAX_PDF_PAGES = 10
MAX_PDF_OCR_PAGES = 3
MAX_EXCEL_ROWS = 100
MAX_PPTX_SLIDES = 25
MAX_ZIP_FILES = 5
MAX_ZIP_TOTAL_BYTES = 10 * 1024 * 1024  # 10 MB uncompressed ceiling
MAX_OUTPUT_CHARS = 25000


def _truncate_output(text: str) -> str:
    """Enforces character ceiling to prevent LLM context window overflows."""
    clean = text.strip()
    if len(clean) > MAX_OUTPUT_CHARS:
        return clean[:MAX_OUTPUT_CHARS] + "\n... [Attachment content truncated for brevity]"
    return clean


def extract_from_pdf(file_bytes: bytes) -> str:
    """
    Extracts text from PDFs with layout awareness, empty-password decryption,
    and automatic Cloud OCR fallback on scanned pages.
    """
    extracted_pages: List[str] = []
    ocr_attempts = 0
    stream = io.BytesIO(file_bytes)

    try:
        reader = PdfReader(stream)

        # Handle empty/null passwords common on auto-generated flight slips & invoices
        if reader.is_encrypted:
            for pwd in ["", b"", " ", b" "]:
                try:
                    if reader.decrypt(pwd) > 0:
                        break
                except Exception:
                    continue

        total_pages = min(len(reader.pages), MAX_PDF_PAGES)

        for idx in range(total_pages):
            page = reader.pages[idx]
            page_text = ""
            try:
                page_text = page.extract_text(extraction_mode="layout") or page.extract_text() or ""
            except Exception as read_err:
                logger.debug(f"Direct text read failed on PDF page {idx}: {read_err}")

            if page_text.strip():
                extracted_pages.append(f"[Page {idx + 1}]\n{page_text.strip()}")
            elif ocr_attempts < MAX_PDF_OCR_PAGES:
                # Scanned page fallback: extract and OCR embedded images
                ocr_fragments: List[str] = []
                try:
                    images = getattr(page, "images", [])
                    for img in images:
                        if ocr_attempts >= MAX_PDF_OCR_PAGES:
                            break
                        ocr_res = extract_text_from_image(img.data)
                        if ocr_res.strip():
                            ocr_fragments.append(ocr_res.strip())
                            ocr_attempts += 1
                except Exception as ocr_err:
                    logger.debug(f"OCR scanning skipped on PDF page {idx}: {ocr_err}")

                if ocr_fragments:
                    extracted_pages.append(f"[Page {idx + 1} - OCR]\n" + "\n".join(ocr_fragments))

    except Exception as e:
        logger.error(f"PDF extraction error: {e}")
        # Secondary fallback: Attempt raw Cloud OCR if document is a flattened single-page scan
        try:
            raw_ocr = extract_text_from_image(file_bytes)
            if raw_ocr.strip():
                return _truncate_output(raw_ocr)
        except Exception:
            pass
    finally:
        stream.close()

    return _truncate_output("\n\n".join(extracted_pages))


def extract_from_docx(file_bytes: bytes) -> str:
    """Extracts paragraphs and tabular data from Word (.docx) files."""
    stream = io.BytesIO(file_bytes)
    try:
        doc = Document(stream)
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        tables_text = []

        for table in doc.tables:
            for row in table.rows:
                cells_seen = set()
                row_cells = []
                for cell in row.cells:
                    ctext = cell.text.strip()
                    # Skip duplicate merged cell text
                    if ctext and ctext not in cells_seen:
                        cells_seen.add(ctext)
                        row_cells.append(ctext)
                if row_cells:
                    tables_text.append(" | ".join(row_cells))

        return _truncate_output("\n".join(paragraphs + tables_text))
    except Exception as e:
        logger.error(f"DOCX extraction error: {e}")
        return ""
    finally:
        stream.close()


def extract_from_excel(file_bytes: bytes, filename: str) -> str:
    """Extracts spreadsheet cells using low-memory streaming readers."""
    rows_text: List[str] = []
    lower = filename.lower()
    stream = io.BytesIO(file_bytes)

    try:
        if (lower.endswith(".xlsx") or lower.endswith(".xlsm")) and openpyxl:
            # read_only=True avoids loading cell styles and formatting trees into RAM
            wb = openpyxl.load_workbook(stream, read_only=True, data_only=True)
            for sheet in wb.worksheets:
                rows_text.append(f"[Sheet: {sheet.title}]")
                row_count = 0
                for row in sheet.iter_rows(values_only=True):
                    if row_count >= MAX_EXCEL_ROWS:
                        rows_text.append("... [Sheet rows truncated]")
                        break
                    filtered = [str(c).strip() for c in row if c is not None and str(c).strip()]
                    if filtered:
                        rows_text.append(" | ".join(filtered))
                        row_count += 1
            wb.close()
            return _truncate_output("\n".join(rows_text))

        elif lower.endswith(".xls") and xlrd:
            wb = xlrd.open_workbook(file_contents=file_bytes)
            for sheet in wb.sheets():
                rows_text.append(f"[Sheet: {sheet.name}]")
                max_r = min(sheet.nrows, MAX_EXCEL_ROWS)
                for row_idx in range(max_r):
                    vals = [str(sheet.cell_value(row_idx, c)).strip() for c in range(sheet.ncols)]
                    filtered = [v for v in vals if v]
                    if filtered:
                        rows_text.append(" | ".join(filtered))
                if sheet.nrows > MAX_EXCEL_ROWS:
                    rows_text.append("... [Sheet rows truncated]")
            return _truncate_output("\n".join(rows_text))

    except Exception as e:
        logger.error(f"Excel extraction error on '{filename}': {e}")
    finally:
        stream.close()

    return ""


def extract_from_pptx(file_bytes: bytes) -> str:
    """Extracts text content from PowerPoint presentations."""
    if not Presentation:
        return ""

    stream = io.BytesIO(file_bytes)
    try:
        prs = Presentation(stream)
        slides_text: List[str] = []
        max_slides = min(len(prs.slides), MAX_PPTX_SLIDES)

        for idx in range(max_slides):
            slide = prs.slides[idx]
            slide_content = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        p_text = paragraph.text.strip()
                        if p_text:
                            slide_content.append(p_text)
            if slide_content:
                slides_text.append(f"[Slide {idx + 1}]\n" + "\n".join(slide_content))

        return _truncate_output("\n\n".join(slides_text))
    except Exception as e:
        logger.error(f"PPTX extraction error: {e}")
        return ""
    finally:
        stream.close()


def extract_from_rtf(file_bytes: bytes) -> str:
    """Extracts plain text from RTF payloads."""
    if not rtf_to_text:
        return ""
    try:
        raw = file_bytes.decode("utf-8", errors="ignore")
        return _truncate_output(rtf_to_text(raw).strip())
    except Exception as e:
        logger.error(f"RTF extraction error: {e}")
        return ""


def extract_from_html(file_bytes: bytes) -> str:
    """Cleans and extracts legible text from HTML attachments."""
    if not BeautifulSoup:
        clean = re.sub(r"<[^>]+>", " ", file_bytes.decode("utf-8", errors="ignore"))
        return _truncate_output(re.sub(r"\s+", " ", clean).strip())
    try:
        soup = BeautifulSoup(file_bytes, "html.parser")
        for tag in soup(["script", "style", "meta", "noscript"]):
            tag.decompose()
        return _truncate_output(soup.get_text(separator="\n", strip=True))
    except Exception as e:
        logger.error(f"HTML extraction error: {e}")
        return ""


def extract_from_ics(file_bytes: bytes) -> str:
    """Parses standard iCalendar (.ics) invite metadata."""
    try:
        cal = icalendar.Calendar.from_ical(file_bytes)
        details = []
        for component in cal.walk():
            if component.name == "VEVENT":
                summary = str(component.get("summary", "")).strip()
                description = str(component.get("description", "")).strip()
                start = component.get("dtstart")
                end = component.get("dtend")
                location = str(component.get("location", "")).strip()

                start_val = start.dt.isoformat() if hasattr(start, "dt") else str(start or "")
                end_val = end.dt.isoformat() if hasattr(end, "dt") else str(end or "")

                details.append(
                    f"Calendar Invite: {summary}\nStart: {start_val}\nEnd: {end_val}\n"
                    f"Location: {location}\nDetails: {description}"
                )
        return _truncate_output("\n---\n".join(details))
    except Exception as e:
        logger.error(f"ICS extraction error: {e}")
        return ""


def extract_from_msg(file_bytes: bytes) -> str:
    """Parses Outlook .msg email archives."""
    if not extract_msg:
        return ""

    stream = io.BytesIO(file_bytes)
    msg = None
    try:
        msg = extract_msg.Message(stream)
        return _truncate_output(
            f"Subject: {msg.subject}\nFrom: {msg.sender}\nDate: {msg.date}\n\n{msg.body or ''}"
        )
    except Exception as e:
        logger.error(f"MSG extraction error: {e}")
        return ""
    finally:
        if msg:
            try:
                msg.close()
            except Exception:
                pass
        stream.close()


def extract_from_eml(file_bytes: bytes) -> str:
    """Parses MIME RFC 822 .eml messages."""
    try:
        msg = email.message_from_bytes(file_bytes)
        body = ""
        if msg.is_multipart():
            for part in msg.walk():
                if part.get_content_type() == "text/plain":
                    body += part.get_payload(decode=True).decode("utf-8", errors="ignore") + "\n"
        else:
            body = msg.get_payload(decode=True).decode("utf-8", errors="ignore")
        return _truncate_output(
            f"Subject: {msg.get('subject', '')}\nDate: {msg.get('date', '')}\n\n{body}"
        )
    except Exception as e:
        logger.error(f"EML extraction error: {e}")
        return ""


def extract_from_zip(file_bytes: bytes) -> str:
    """Extracts text from files within zip archives with decompression bomb protection."""
    extracted_docs = []
    total_uncompressed = 0
    stream = io.BytesIO(file_bytes)

    try:
        with zipfile.ZipFile(stream) as z:
            count = 0
            for info in z.infolist():
                fname = info.filename
                if fname.startswith("__MACOSX") or fname.endswith("/") or count >= MAX_ZIP_FILES:
                    continue

                total_uncompressed += info.file_size
                if total_uncompressed > MAX_ZIP_TOTAL_BYTES:
                    extracted_docs.append("[Archive uncompressed size limit reached]")
                    break

                nested_bytes = z.read(fname)
                nested_text = extract_attachment_text(fname, nested_bytes)
                if nested_text.strip():
                    extracted_docs.append(f"--- Archive Item: {fname} ---\n{nested_text.strip()}")
                    count += 1

        return _truncate_output("\n\n".join(extracted_docs))
    except Exception as e:
        logger.error(f"ZIP extraction error: {e}")
        return ""
    finally:
        stream.close()


def extract_attachment_text(filename: str, file_bytes: bytes, mime_type: str = "") -> str:
    """Universal attachment router dispatching to file-specific extractors."""
    if not filename or not file_bytes:
        return ""

    lower = filename.lower()

    if lower.endswith(".pdf"):
        return extract_from_pdf(file_bytes)
    elif lower.endswith(".docx"):
        return extract_from_docx(file_bytes)
    elif lower.endswith(".rtf"):
        return extract_from_rtf(file_bytes)
    elif lower.endswith((".xlsx", ".xlsm", ".xls")):
        return extract_from_excel(file_bytes, filename)
    elif lower.endswith((".pptx", ".ppt")):
        return extract_from_pptx(file_bytes)
    elif lower.endswith((".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tiff", ".tif")):
        return _truncate_output(extract_text_from_image(file_bytes))
    elif lower.endswith(".ics"):
        return extract_from_ics(file_bytes)
    elif lower.endswith(".msg"):
        return extract_from_msg(file_bytes)
    elif lower.endswith(".eml"):
        return extract_from_eml(file_bytes)
    elif lower.endswith((".html", ".htm", ".xhtml", ".xml")):
        return extract_from_html(file_bytes)
    elif lower.endswith((".zip", ".jar")):
        return extract_from_zip(file_bytes)
    elif lower.endswith((".txt", ".csv", ".tsv", ".json", ".md", ".yaml", ".yml", ".log")):
        try:
            return _truncate_output(file_bytes.decode("utf-8", errors="ignore"))
        except Exception:
            return _truncate_output(file_bytes.decode("latin-1", errors="ignore"))

    # Plain-text inspection heuristic
    if b"\x00" not in file_bytes[:1024]:
        try:
            decoded = file_bytes.decode("utf-8", errors="ignore").strip()
            if len(decoded) > 10:
                return _truncate_output(decoded)
        except Exception:
            pass

    return ""